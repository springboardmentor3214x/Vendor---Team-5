"""Rendering-only Module 10 exports; report aggregation remains in report_service."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table as PdfTable, TableStyle

SYSTEM_NAME = "Vendor Reliability Intelligence Platform"
REPORT_TITLES = {
    "vendor_performance": "Vendor Performance Report", "procurement_summary": "Procurement Report",
    "purchase_order": "Purchase Order Report", "compliance": "Compliance Report",
    "contract": "Contract Report", "executive_summary": "Executive Summary Report",
}


def _value(value: Any) -> str:
    return "" if value is None else str(value)


def _columns(rows: list[dict[str, Any]]) -> list[str]:
    return list(dict.fromkeys(key for row in rows for key in row)) or ["message"]


def _normalise_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return rows or [{"message": "No matching records"}]


def build_excel(report_type: str, rows: list[dict[str, Any]], filters: dict[str, Any]) -> bytes:
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Report Data"
    data, columns = _normalise_rows(rows), _columns(rows)
    sheet.append(columns)
    for row in data: sheet.append([_value(row.get(column)) for column in columns])
    for cell in sheet[1]: cell.font = Font(bold=True)
    table = Table(displayName="ReportData", ref=f"A1:{chr(64 + len(columns))}{len(data) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    metadata = workbook.create_sheet("Export Details")
    metadata.append(["Report", REPORT_TITLES[report_type]])
    metadata.append(["Generated at", datetime.utcnow().isoformat()])
    metadata.append(["Applied filters", "No filters applied" if not filters else ""])
    for key, value in filters.items(): metadata.append([key, _value(value)])
    output = BytesIO(); workbook.save(output); return output.getvalue()


def _footer(canvas, document):
    canvas.saveState(); canvas.setFont("Helvetica", 8)
    canvas.drawString(36, 20, f"{SYSTEM_NAME} — {datetime.utcnow():%Y-%m-%d}")
    canvas.drawRightString(A4[1] - 36, 20, f"Page {document.page}")
    canvas.restoreState()


def build_pdf(report_type: str, rows: list[dict[str, Any]], filters: dict[str, Any], chart_data: dict[str, Any] | None = None) -> bytes:
    output = BytesIO(); doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=32, bottomMargin=34)
    styles = getSampleStyleSheet(); story = [Paragraph(SYSTEM_NAME, styles["Heading2"]), Paragraph(REPORT_TITLES[report_type], styles["Title"]), Paragraph(f"Generated: {datetime.utcnow():%Y-%m-%d %H:%M UTC}", styles["Normal"])]
    story += [Spacer(1, 6), Paragraph("Applied filters: " + ("; ".join(f"{k}={v}" for k, v in filters.items()) if filters else "No filters applied"), styles["Normal"])]
    data, columns = _normalise_rows(rows), _columns(rows)
    story += [Spacer(1, 10), PdfTable([columns] + [[_value(row.get(column))[:80] for column in columns] for row in data], repeatRows=1, style=TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#163b65")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .25, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 7)]))]
    series = (chart_data or {}).get("series") or (chart_data or {}).get("spending_by_vendor") or []
    if series and isinstance(series, list) and isinstance(series[0], dict):
        labels = [str(item.get("label") or item.get("vendor") or item.get("month") or "") for item in series[:8]]
        values = [float(item.get("total") or item.get("count") or 0) for item in series[:8]]
        drawing = Drawing(420, 170); chart = VerticalBarChart(); chart.data = [values]; chart.categoryAxis.categoryNames = labels; chart.width, chart.height = 360, 120; chart.x, chart.y = 35, 30; drawing.add(chart); story += [Spacer(1, 12), drawing]
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer); return output.getvalue()
