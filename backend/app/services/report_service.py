"""Module 10 data-driven reports, chart shaping, and export renderers."""
from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import date, datetime
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.compliance import ComplianceRecord
from app.models.contract import Contract
from app.models.invoice import Invoice
from app.models.performance import PerformanceRecord
from app.models.procurement_request import ProcurementRequest
from app.models.purchase_order import PurchaseOrder
from app.models.reliability import VendorReliability
from app.models.vendor import Vendor
from app.models.vendor_document import VendorDocument
from app.services.notification_service import get_user_notifications

_CONTROL_FILTERS = {"sort_by", "sort_order", "order", "limit", "offset", "start_date", "end_date"}

def _value(value: Any) -> Any:
    return value.isoformat() if isinstance(value, (date, datetime)) else value

def _rows(db: Any, model: Any) -> list[Any]:
    return list(db.query(model).all() or []) if db is not None else []

def _matches(item: Any, filters: Optional[Dict[str, Any]]) -> bool:
    if not filters:
        return True
    for key, expected in filters.items():
        if key in _CONTROL_FILTERS or expected is None or not hasattr(item, key):
            continue
        actual = getattr(item, key)
        if isinstance(expected, (list, tuple, set)):
            if actual not in expected:
                return False
        elif str(actual).lower() != str(expected).lower():
            return False
    created = getattr(item, "created_at", None) or getattr(item, "request_date", None) or getattr(item, "po_date", None)
    start, end = (filters or {}).get("start_date"), (filters or {}).get("end_date")
    created_day = created.date() if isinstance(created, datetime) else created
    return not ((start and created_day and created_day < start) or (end and created_day and created_day > end))

def _filtered(db: Any, model: Any, filters: Optional[Dict[str, Any]] = None) -> list[Any]:
    return [row for row in _rows(db, model) if _matches(row, filters)]

def _shape(rows: list[dict[str, Any]], filters: Optional[Dict[str, Any]] = None) -> list[dict[str, Any]]:
    filters = filters or {}
    key, reverse = filters.get("sort_by"), str(filters.get("sort_order", filters.get("order", "asc"))).lower() == "desc"
    if key:
        rows.sort(key=lambda row: (row.get(key) is None, str(row.get(key)).lower()), reverse=reverse)
    offset, limit = int(filters.get("offset", 0) or 0), filters.get("limit")
    return rows[offset: offset + int(limit)] if limit is not None else rows[offset:]

def generate_purchase_order_report(db: Any, filters: Optional[Dict[str, Any]] = None) -> list[dict[str, Any]]:
    vendors = {v.id: v for v in _rows(db, Vendor)}
    invoices = {i.purchase_order_id: i for i in _rows(db, Invoice)}
    result = []
    for po in _filtered(db, PurchaseOrder, filters):
        vendor, invoice = vendors.get(getattr(po, "vendor_id", None)), invoices.get(getattr(po, "id", None))
        result.append({"purchase_order_id": po.id, "purchase_order_number": getattr(po, "po_number", None),
                       "vendor_id": getattr(po, "vendor_id", None), "vendor_name": getattr(vendor, "company_name", None),
                       "purchase_date": _value(getattr(po, "po_date", None) or getattr(po, "created_at", None)),
                       "expected_delivery_date": _value(getattr(po, "expected_delivery_date", None)),
                       "delivery_date": _value(getattr(po, "actual_delivery_date", None)),
                       "order_value": float(getattr(po, "total_cost", 0) or 0),
                       "current_status": getattr(po, "po_status", None), "invoice_status": getattr(invoice, "payment_status", None)})
    return _shape(result, filters)

def generate_contract_report(db: Any, filters: Optional[Dict[str, Any]] = None) -> list[dict[str, Any]]:
    vendors, today = {v.id: v for v in _rows(db, Vendor)}, date.today()
    result = []
    for contract in _filtered(db, Contract, filters):
        end = getattr(contract, "end_date", None); end = end.date() if isinstance(end, datetime) else end
        result.append({"id": contract.id, "vendor_id": contract.vendor_id, "vendor_name": getattr(vendors.get(contract.vendor_id), "company_name", None),
                       "contract_number": getattr(contract, "contract_number", None), "contract_title": getattr(contract, "contract_title", None),
                       "status": getattr(contract, "status", None), "start_date": _value(getattr(contract, "start_date", None)), "end_date": _value(end),
                       "contract_value": float(getattr(contract, "contract_value", 0) or 0), "responsible_manager": getattr(contract, "responsible_manager", None),
                       "days_to_expiry": (end - today).days if end else None})
    return _shape(result, filters)

def generate_compliance_report(db: Any, filters: Optional[Dict[str, Any]] = None) -> list[dict[str, Any]]:
    return _shape([{"id": r.id, "vendor_id": r.vendor_id, "compliance_type": getattr(r, "compliance_type", None), "status": getattr(r, "status", None), "verification_date": _value(getattr(r, "verification_date", None)), "remarks": getattr(r, "remarks", None)} for r in _filtered(db, ComplianceRecord, filters)], filters)

def generate_vendor_document_report(db: Any, filters: Optional[Dict[str, Any]] = None) -> list[dict[str, Any]]:
    return _shape([{key: _value(getattr(r, key, None)) for key in ("id", "vendor_id", "document_type", "file_name", "content_type", "uploaded_at")} for r in _filtered(db, VendorDocument, filters)], filters)

def generate_vendor_performance_report(db: Any, filters: Optional[Dict[str, Any]] = None) -> list[dict[str, Any]]:
    reliability = {r.vendor_id: r for r in _rows(db, VendorReliability)}
    return _shape([{**{key: _value(getattr(r, key, None)) for key in ("id", "vendor_id", "total_completed_orders", "on_time_delivery_rate", "average_quality_score", "overall_performance_score", "performance_status", "evaluation_date")}, "reliability_score": _value(getattr(reliability.get(r.vendor_id), "reliability_score", None))} for r in _filtered(db, PerformanceRecord, filters)], filters)

def generate_procurement_summary_report(db: Any, filters: Optional[Dict[str, Any]] = None) -> list[dict[str, Any]]:
    orders = {o.procurement_request_id: o for o in _rows(db, PurchaseOrder)}
    return _shape([{ "procurement_request_id": r.id, "request_number": getattr(r, "request_number", None), "title": getattr(r, "title", None), "department": getattr(r, "department", None), "project_name": getattr(r, "project_name", None), "vendor_id": getattr(r, "vendor_id", None), "estimated_budget": getattr(r, "estimated_budget", None), "approval_status": getattr(r, "approval_status", None), "po_number": getattr(orders.get(r.id), "po_number", None), "total_cost": getattr(orders.get(r.id), "total_cost", None)} for r in _filtered(db, ProcurementRequest, filters)], filters)

def generate_notification_report(db: Any, filters: Optional[Dict[str, Any]] = None) -> list[dict[str, Any]]:
    user_id = (filters or {}).get("user_id")
    if user_id is None: return []
    return _shape([{key: _value(getattr(row, key, None)) for key in ("id", "user_id", "title", "message", "notification_type", "related_module", "priority", "is_read", "created_at")} for row in get_user_notifications(db, user_id, filters=filters)], filters)

def generate_executive_summary_report(db: Any, filters: Optional[Dict[str, Any]] = None) -> dict[str, Any]:
    pos, contracts, vendors = generate_purchase_order_report(db, filters), generate_contract_report(db, filters), _filtered(db, Vendor, filters)
    compliance_records = _filtered(db, ComplianceRecord, filters)
    compliance_percentage = None
    if compliance_records:
        compliance_percentage = round(
            100 * sum(str(getattr(record, "status", "")).casefold() == "compliant" for record in compliance_records) / len(compliance_records),
            2,
        )
    return {"total_registered_vendors": len(vendors), "total_procurement_expenditure": sum(row["order_value"] for row in pos), "completed_purchase_orders": sum(str(row["current_status"]).lower() in {"completed", "delivered"} for row in pos), "contracts_near_expiry_count": sum(0 <= (row["days_to_expiry"] or -1) <= 30 for row in contracts), "compliance_percentage": compliance_percentage}

def shape_chart_data(rows: list[dict[str, Any]], label_field: str, value_field: str, chart_type: str = "bar") -> dict[str, Any]:
    totals: dict[str, float] = defaultdict(float)
    for row in rows: totals[str(row.get(label_field) or "Unspecified")] += float(row.get(value_field) or 0)
    return {"chart_type": chart_type, "labels": list(totals), "datasets": [{"label": value_field.replace("_", " ").title(), "data": list(totals.values())}]}

_REPORT_CHART_FIELDS = {
    "vendor-performance": ("vendor_id", "overall_performance_score"), "vendor_performance": ("vendor_id", "overall_performance_score"),
    "procurement": ("department", "estimated_budget"), "procurement_summary": ("department", "estimated_budget"),
    "purchase-orders": ("purchase_order_number", "order_value"), "purchase_order": ("purchase_order_number", "order_value"),
    "compliance": ("status", "id"), "contract": ("contract_number", "contract_value"), "contracts": ("contract_number", "contract_value"),
    "vendor_document": ("document_type", "id"), "notification": ("notification_type", "id"),
    "executive_summary": ("total_registered_vendors", "total_procurement_expenditure"),
}

def export_report_data(db: Any, report_type: str, format: str = "csv", filters: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    generators = {"vendor-performance": generate_vendor_performance_report, "vendor_performance": generate_vendor_performance_report, "procurement": generate_procurement_summary_report, "procurement_summary": generate_procurement_summary_report, "purchase-orders": generate_purchase_order_report, "purchase_order": generate_purchase_order_report, "compliance": generate_compliance_report, "contract": generate_contract_report, "contracts": generate_contract_report, "vendor_document": generate_vendor_document_report, "notification": generate_notification_report, "executive_summary": generate_executive_summary_report}
    if report_type not in generators: raise ValueError(f"Unsupported report type: {report_type}")
    if format.lower() not in {"csv", "pdf", "excel"}: raise ValueError(f"Unsupported export format: {format}")
    rows = generators[report_type](db, filters); rows = rows if isinstance(rows, list) else [rows]
    default_label, default_value = _REPORT_CHART_FIELDS[report_type]
    return {"report_type": report_type, "format": format.lower(), "rows": rows, "chart_data": shape_chart_data(rows, (filters or {}).get("chart_label", default_label), (filters or {}).get("chart_value", default_value), (filters or {}).get("chart_type", "bar")), "metadata": {"title": f"{report_type.replace('-', ' ').title()} Report", "generated_at": datetime.utcnow().isoformat(), "row_count": len(rows), "status": "prepared"}}

def render_excel_csv_report(rows: List[Dict[str, Any]]) -> str:
    output = io.StringIO(); fields = list(rows[0]) if rows else ["status_message"]
    writer = csv.DictWriter(output, fieldnames=fields); writer.writeheader(); writer.writerows(rows or [{"status_message": "No data matching applied filters"}]); return output.getvalue()

def render_excel_report(rows: List[Dict[str, Any]], title: str = "Report") -> bytes:
    workbook = Workbook(); sheet = workbook.active; sheet.title = title[:31]
    fields = list(rows[0]) if rows else ["status_message"]; sheet.append(fields)
    for cell in sheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows or [{"status_message": "No data matching applied filters"}]: sheet.append([_value(row.get(field)) for field in fields])
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    for index, field in enumerate(fields, 1): sheet.column_dimensions[get_column_letter(index)].width = min(max(len(field) + 2, 14), 35)
    output = io.BytesIO(); workbook.save(output); return output.getvalue()

def render_pdf_report(report_title: str, rows: List[Dict[str, Any]], metadata: Optional[Dict[str, Any]] = None) -> bytes:
    output = io.BytesIO(); styles = getSampleStyleSheet(); story = [Paragraph("Vendor Reliability Intelligence Platform", styles["Title"]), Paragraph(report_title, styles["Heading2"]), Paragraph(f"Generated {datetime.utcnow():%Y-%m-%d %H:%M UTC} | {len(rows)} records", styles["Normal"]), Spacer(1, 0.2 * inch)]
    fields = list(rows[0]) if rows else ["status_message"]; data = [fields] + [[str(_value(row.get(field)) or "")[:45] for field in fields] for row in (rows or [{"status_message": "No data matching applied filters"}])]
    table = Table(data, repeatRows=1); table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1F4E78")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("GRID", (0,0), (-1,-1), .25, colors.grey), ("VALIGN", (0,0), (-1,-1), "TOP"), ("FONTSIZE", (0,0), (-1,-1), 7)])); story.append(table)
    SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=28).build(story); return output.getvalue()
