from io import BytesIO

from openpyxl import load_workbook

from app.services.export_service import build_excel, build_pdf


REPORT_TYPES = ["vendor_performance", "procurement_summary", "purchase_order", "compliance", "contract", "executive_summary"]


def test_pdf_exports_for_required_report_types() -> None:
    for report_type in REPORT_TYPES:
        payload = build_pdf(report_type, [{"vendor_name": "Demo Vendor", "score": 90}], {"department": "IT"}, {"series": []})
        assert payload.startswith(b"%PDF-")


def test_excel_exports_for_required_report_types() -> None:
    for report_type in REPORT_TYPES:
        payload = build_excel(report_type, [{"vendor_name": "Demo Vendor", "score": 90}], {"department": "IT"})
        workbook = load_workbook(BytesIO(payload))
        assert workbook["Report Data"]["A1"].value == "vendor_name"
        assert "ReportData" in workbook["Report Data"].tables


def test_pdf_without_chart_data_is_valid() -> None:
    assert build_pdf("contract", [{"contract_number": "C-1"}], {}, {"series": [], "reason": "No chart"}).startswith(b"%PDF-")
