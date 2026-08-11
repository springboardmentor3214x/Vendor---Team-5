"""HTTP-level Module 10 PDF/XLSX export regression tests."""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.main import app


client = TestClient(app)
REPORT_KEYS = ["vendor-performance", "procurement", "purchase-orders", "compliance", "contracts", "executive-summary"]
FILTER_VALUE = "Module10ExportProof"


@pytest.fixture(scope="module")
def report_headers() -> dict[str, str]:
    email, password = "module10.exports@example.com", "Password123"
    response = client.post("/auth/register", json={
        "fullName": "Module 10 Export Admin", "email": email, "password": password,
        "confirmPassword": password, "role": "Administrator",
    })
    assert response.status_code in {201, 400}, response.text
    login = client.post("/auth/login", json={"email": email, "password": password})
    assert login.status_code == 200, login.text
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


@pytest.mark.parametrize("report_key", REPORT_KEYS)
def test_filtered_pdf_export_http(report_key: str, report_headers: dict[str, str]) -> None:
    response = client.get(f"/reports/{report_key}/export", params={"format": "pdf", "department": FILTER_VALUE}, headers=report_headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/pdf")
    assert f'attachment; filename="{report_key.replace("-", "_")}.pdf"' == response.headers["content-disposition"]
    assert response.content.startswith(b"%PDF-")
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(response.content)).pages)
    assert FILTER_VALUE in text  # Actual PDF text extraction, not status inference.


@pytest.mark.parametrize("report_key", REPORT_KEYS)
def test_filtered_xlsx_export_http(report_key: str, report_headers: dict[str, str]) -> None:
    response = client.get(f"/reports/{report_key}/export", params={"format": "xlsx", "department": FILTER_VALUE}, headers=report_headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert f'attachment; filename="{report_key.replace("-", "_")}.xlsx"' == response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content))
    values = [cell.value for row in workbook["Export Details"].iter_rows() for cell in row]
    assert FILTER_VALUE in values  # Actual downloaded workbook metadata.


def test_export_endpoint_requires_jwt() -> None:
    response = client.get("/reports/vendor-performance/export", params={"format": "pdf", "department": FILTER_VALUE})
    assert response.status_code == 401
